from __future__ import annotations

import csv
from collections import OrderedDict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_FILE = Path("/Users/aniketshinde/Downloads/ROUTE EXCEL.xlsx")
OUTPUT_DIR = (
    Path("/Users/aniketshinde/Projects/frappe-bench/apps/trans_ms")
    / "output"
    / "spreadsheet"
    / "erpnext_v15_trip_route_imports"
)

TRIP_ROUTE_HEADERS = [
    "Route Name",
    "Total Distance (km)",
    "Location (trip_steps)",
    "Distance (trip_steps)",
    "Location Type (trip_steps)",
    "Fuel Consumption Qty (Ltr.) (trip_steps)",
]

TRIP_LOCATION_HEADERS = [
    "Description",
    "Location (Latitude, Longitude)",
    "Is Local Border",
    "Is International Border",
]

TRIP_LOCATION_TYPE_HEADERS = [
    "Location Type",
    "Arrival Date",
    "Departure Date",
    "Loading Date",
    "Offloading Date",
]

ROUTE_MAP_HEADERS = [
    "source_row",
    "route_name",
    "route_from",
    "route_to",
    "status",
    "notes",
]

REVIEW_HEADERS = [
    "level",
    "source_row",
    "route_name",
    "issue",
    "details",
]

DEFAULT_LOCATION_TYPES = [
    {
        "Location Type": "Loading Point",
        "Arrival Date": 0,
        "Departure Date": 0,
        "Loading Date": 1,
        "Offloading Date": 0,
    },
    {
        "Location Type": "Offloading Point",
        "Arrival Date": 0,
        "Departure Date": 0,
        "Loading Date": 0,
        "Offloading Date": 1,
    },
    {
        "Location Type": "Border",
        "Arrival Date": 1,
        "Departure Date": 1,
        "Loading Date": 0,
        "Offloading Date": 0,
    },
]


def clean_value(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def write_csv(path: Path, headers: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in headers})


def write_xlsx(path: Path, sheet_name: str, headers: list[str], rows: list[dict[str, object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for index, header in enumerate(headers, start=1):
        max_length = len(header)
        for row_index in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row=row_index, column=index).value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 80)

    workbook.save(path)


def write_dataset(base_path: Path, sheet_name: str, headers: list[str], rows: list[dict[str, object]]) -> None:
    write_csv(base_path.with_suffix(".csv"), headers, rows)
    write_xlsx(base_path.with_suffix(".xlsx"), sheet_name, headers, rows)


def read_source_rows() -> list[dict[str, object]]:
    workbook = load_workbook(SOURCE_FILE, data_only=True)
    worksheet = workbook.active
    rows: list[dict[str, object]] = []
    for row_number in range(2, worksheet.max_row + 1):
        route_key = clean_value(worksheet.cell(row_number, 1).value)
        route_from = clean_value(worksheet.cell(row_number, 2).value)
        route_to = clean_value(worksheet.cell(row_number, 3).value)
        if not (route_key or route_from or route_to):
            continue
        rows.append(
            {
                "source_row": row_number,
                "route_name": route_key,
                "route_from": route_from,
                "route_to": route_to,
            }
        )
    return rows


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    source_rows = read_source_rows()

    route_import_rows: list[dict[str, object]] = []
    trip_location_rows_map: OrderedDict[str, dict[str, object]] = OrderedDict()
    route_map_rows: list[dict[str, object]] = []
    review_rows: list[dict[str, object]] = []

    seen_route_pairs: set[tuple[str, str, str]] = set()
    kept_route_count = 0
    duplicate_count = 0
    blank_destination_count = 0
    same_origin_destination_count = 0

    for row in source_rows:
        route_name = clean_value(row["route_name"])
        route_from = clean_value(row["route_from"])
        route_to = clean_value(row["route_to"])
        source_row = row["source_row"]

        if not route_name:
            review_rows.append(
                {
                    "level": "error",
                    "source_row": source_row,
                    "route_name": "",
                    "issue": "Missing route name",
                    "details": "Skipped because `Trip Route` requires Route Name.",
                }
            )
            route_map_rows.append(
                {
                    "source_row": source_row,
                    "route_name": "",
                    "route_from": route_from,
                    "route_to": route_to,
                    "status": "skipped",
                    "notes": "Missing RouteKey",
                }
            )
            continue

        if not route_from or not route_to:
            blank_destination_count += 1
            review_rows.append(
                {
                    "level": "warning",
                    "source_row": source_row,
                    "route_name": route_name,
                    "issue": "Missing route endpoint",
                    "details": "Skipped because both RouteFrom and RouteTo are required to build a two-step route.",
                }
            )
            route_map_rows.append(
                {
                    "source_row": source_row,
                    "route_name": route_name,
                    "route_from": route_from,
                    "route_to": route_to,
                    "status": "skipped",
                    "notes": "Missing RouteFrom or RouteTo",
                }
            )
            continue

        route_signature = (route_name, route_from, route_to)
        if route_signature in seen_route_pairs:
            duplicate_count += 1
            review_rows.append(
                {
                    "level": "info",
                    "source_row": source_row,
                    "route_name": route_name,
                    "issue": "Duplicate route row",
                    "details": "Skipped exact duplicate of an earlier row with the same RouteKey, RouteFrom, and RouteTo.",
                }
            )
            route_map_rows.append(
                {
                    "source_row": source_row,
                    "route_name": route_name,
                    "route_from": route_from,
                    "route_to": route_to,
                    "status": "deduped",
                    "notes": "Exact duplicate skipped",
                }
            )
            continue

        seen_route_pairs.add(route_signature)
        kept_route_count += 1

        if route_from == route_to:
            same_origin_destination_count += 1
            review_rows.append(
                {
                    "level": "info",
                    "source_row": source_row,
                    "route_name": route_name,
                    "issue": "Origin and destination are the same",
                    "details": "Imported as provided, but this route may need business review.",
                }
            )

        route_import_rows.append(
            {
                "Route Name": route_name,
                "Total Distance (km)": 0,
                "Location (trip_steps)": route_from,
                "Distance (trip_steps)": "",
                "Location Type (trip_steps)": "Loading Point",
                "Fuel Consumption Qty (Ltr.) (trip_steps)": "",
            }
        )
        route_import_rows.append(
            {
                "Route Name": "",
                "Total Distance (km)": "",
                "Location (trip_steps)": route_to,
                "Distance (trip_steps)": "",
                "Location Type (trip_steps)": "Offloading Point",
                "Fuel Consumption Qty (Ltr.) (trip_steps)": "",
            }
        )

        for location in (route_from, route_to):
            if location not in trip_location_rows_map:
                trip_location_rows_map[location] = {
                    "Description": location,
                    "Location (Latitude, Longitude)": "",
                    "Is Local Border": 0,
                    "Is International Border": 0,
                }

        route_map_rows.append(
            {
                "source_row": source_row,
                "route_name": route_name,
                "route_from": route_from,
                "route_to": route_to,
                "status": "imported",
                "notes": "Total Distance (km) set to 0 because source file does not include distance data.",
            }
        )

    trip_location_rows = list(trip_location_rows_map.values())

    if kept_route_count:
        review_rows.insert(
            0,
            {
                "level": "warning",
                "source_row": "",
                "route_name": "",
                "issue": "Missing distance data in source",
                "details": (
                    "All generated Trip Route rows use `Total Distance (km) = 0` because the workbook only "
                    "contains RouteKey, RouteFrom, and RouteTo. Update distances after import if they are needed "
                    "for downstream fuel or costing logic."
                ),
            }
        )

    write_dataset(OUTPUT_DIR / "trip_location_type_seed_v15", "Trip Location Type", TRIP_LOCATION_TYPE_HEADERS, DEFAULT_LOCATION_TYPES)
    write_dataset(OUTPUT_DIR / "trip_location_import_v15", "Trip Location", TRIP_LOCATION_HEADERS, trip_location_rows)
    write_dataset(OUTPUT_DIR / "trip_route_import_v15", "Trip Route", TRIP_ROUTE_HEADERS, route_import_rows)
    write_dataset(OUTPUT_DIR / "route_source_map_v15", "Route Map", ROUTE_MAP_HEADERS, route_map_rows)
    write_dataset(OUTPUT_DIR / "route_review_v15", "Review", REVIEW_HEADERS, review_rows)

    summary_lines = [
        "# ERPNext v15 Trip Route Import Outputs",
        "",
        f"Source file: `{SOURCE_FILE}`",
        f"Generated on: `{date.today().isoformat()}`",
        "",
        "Matched to the transport route doctypes in this checkout of the transport app:",
        "- `Trip Route`",
        "- `Trip Steps Table`",
        "- `Trip Location`",
        "- `Trip Location Type`",
        "",
        "Import order:",
        "1. `trip_location_type_seed_v15.xlsx` or `.csv`",
        "2. `trip_location_import_v15.xlsx` or `.csv`",
        "3. `trip_route_import_v15.xlsx` or `.csv`",
        "",
        "Counts:",
        f"- Source rows read: {len(source_rows)}",
        f"- Routes generated: {kept_route_count}",
        f"- Trip Route import rows: {len(route_import_rows)}",
        f"- Trip Location rows: {len(trip_location_rows)}",
        f"- Exact duplicates skipped: {duplicate_count}",
        f"- Incomplete rows skipped: {blank_destination_count}",
        f"- Same-origin/destination routes flagged: {same_origin_destination_count}",
        f"- Review rows: {len(review_rows)}",
        "",
        "Important assumptions:",
        "- Each source row was converted into a two-step route.",
        "- Step 1 is `Loading Point` using `RouteFrom`.",
        "- Step 2 is `Offloading Point` using `RouteTo`.",
        "- `Total Distance (km)` was set to `0` for every route because the source file does not provide distance data.",
        "- Child-step distances and fuel consumption quantities were left blank for the same reason.",
        "",
        "Use `route_source_map_v15` to trace each generated route back to the original spreadsheet row.",
    ]
    (OUTPUT_DIR / "README.md").write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    print(f"Generated files in {OUTPUT_DIR}")
    print(
        "Summary:",
        {
            "source_rows": len(source_rows),
            "routes_generated": kept_route_count,
            "trip_route_rows": len(route_import_rows),
            "trip_location_rows": len(trip_location_rows),
            "duplicates_skipped": duplicate_count,
            "incomplete_skipped": blank_destination_count,
            "review_rows": len(review_rows),
        },
    )


if __name__ == "__main__":
    main()

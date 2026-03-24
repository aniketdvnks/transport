from __future__ import annotations

import csv
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_FILE = Path("/Users/aniketshinde/Downloads/PARTY NAME.xlsx")
OUTPUT_DIR = (
    Path("/Users/aniketshinde/Projects/frappe-bench/apps/trans_ms")
    / "output"
    / "spreadsheet"
    / "erpnext_v15_party_imports"
)

CUSTOMER_ROLES = {"CLIENT"}
SUPPLIER_ROLES = {
    "Supplier",
    "Tyre Supplier",
    "PUMP",
    "Tax Vendor",
    "OWNER",
    "Broker",
    "Agent",
    "Financier",
    "Bank",
}
DRIVER_ROLES = {"Driver Ledger", "Debit Driver"}

ACCOUNT_LIKE_ROLES = {
    "Income Head",
    "Trip Expense Ledger",
    "Expense",
    "Account Group",
    "Tax Receivable",
    "Interestaccount",
    "Loanaccount",
    "Deduction Ledger",
    "Hire Purchase Account",
    "GST Ledger",
    "TDS Payable",
    "Tax Payable",
    "Sales Account",
    "Cash/Bank",
    "Debit Driver",
}

COMPANY_KEYWORDS = {
    "LTD",
    "LIMITED",
    "PVT",
    "PRIVATE",
    "LLP",
    "ROADWAYS",
    "TRANSPORT",
    "LOGISTICS",
    "CORPORATION",
    "COMPANY",
    "CO",
    "ENTERPRISE",
    "ENTERPRISES",
    "INDUSTRIES",
    "CHEMICALS",
    "WORKS",
    "SERVICES",
    "MOTORS",
    "FUELS",
    "BANK",
    "FINANCE",
    "INSURANCE",
    "TRAVELS",
    "IMPEX",
    "AGRO",
    "STATION",
}

TRANSPORTER_KEYWORDS = {
    "TRANSPORT",
    "ROADWAYS",
    "LOGISTICS",
    "CARRIERS",
    "FREIGHT",
    "TRANS",
    "TRAVELS",
}

GST_STATE_BY_CODE = {
    "01": "Jammu and Kashmir",
    "02": "Himachal Pradesh",
    "03": "Punjab",
    "04": "Chandigarh",
    "05": "Uttarakhand",
    "06": "Haryana",
    "07": "Delhi",
    "08": "Rajasthan",
    "09": "Uttar Pradesh",
    "10": "Bihar",
    "11": "Sikkim",
    "12": "Arunachal Pradesh",
    "13": "Nagaland",
    "14": "Manipur",
    "15": "Mizoram",
    "16": "Tripura",
    "17": "Meghalaya",
    "18": "Assam",
    "19": "West Bengal",
    "20": "Jharkhand",
    "21": "Odisha",
    "22": "Chhattisgarh",
    "23": "Madhya Pradesh",
    "24": "Gujarat",
    "26": "Dadra and Nagar Haveli and Daman and Diu",
    "27": "Maharashtra",
    "29": "Karnataka",
    "30": "Goa",
    "31": "Lakshadweep Islands",
    "32": "Kerala",
    "33": "Tamil Nadu",
    "34": "Puducherry",
    "35": "Andaman and Nicobar Islands",
    "36": "Telangana",
    "37": "Andhra Pradesh",
    "38": "Ladakh",
    "96": "Other Countries",
    "97": "Other Territory",
}

STATE_NAMES = set(GST_STATE_BY_CODE.values())


def clean_value(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    text = re.sub(r"\s+", " ", text)
    if text.lower() == "none":
        return ""
    return text


def clean_key(text: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "", clean_value(text).upper())


def tidy_address_part(value: str) -> str:
    value = clean_value(value)
    value = re.sub(r"\s*,\s*", ", ", value)
    value = re.sub(r"(,\s*){2,}", ", ", value)
    return value.strip(" ,")


def looks_like_person(name: str) -> bool:
    name = clean_value(name)
    if not name:
        return False
    upper = name.upper()
    if any(keyword in upper for keyword in COMPANY_KEYWORDS):
        return False
    if any(char.isdigit() for char in upper):
        return False
    words = [word for word in re.split(r"[^A-Z]+", upper) if word]
    return 2 <= len(words) <= 6


def infer_party_type(pan: str, name: str, role: str) -> str:
    pan = clean_value(pan).upper()
    if len(pan) >= 4:
        entity_code = pan[3]
        if entity_code == "P":
            return "Individual"
        if entity_code == "F":
            return "Partnership"
        return "Company"

    if role in DRIVER_ROLES or role in {"Employ"}:
        return "Individual"

    if role in {"Broker", "Agent"}:
        return "Individual" if looks_like_person(name) else "Company"

    if role in CUSTOMER_ROLES or role in SUPPLIER_ROLES:
        return "Company"

    return "Individual" if looks_like_person(name) else "Company"


def infer_is_transporter(role: str, name: str) -> int:
    upper = clean_value(name).upper()
    if role == "OWNER":
        return 1
    if any(keyword in upper for keyword in TRANSPORTER_KEYWORDS):
        return 1
    return 0


def parse_place_name(place_name: str) -> str:
    place_name = clean_value(place_name)
    if not place_name:
        return ""

    if "-" in place_name:
        parts = [clean_value(part) for part in place_name.split("-") if clean_value(part)]
        if parts:
            tail = parts[-1]
            if clean_key(tail) not in {clean_key(state) for state in STATE_NAMES}:
                return tail

    return place_name


def extract_pincode(row: dict[str, str]) -> str:
    for field in ("Address1", "Address2", "Address3", "Address4"):
        match = re.search(r"\b(\d{6})\b", clean_value(row.get(field)))
        if match:
            return match.group(1)
    return ""


def infer_state(row: dict[str, str]) -> str:
    gst_no = clean_value(row.get("GSTNo")).upper()
    if len(gst_no) >= 2 and gst_no[:2] in GST_STATE_BY_CODE:
        return GST_STATE_BY_CODE[gst_no[:2]]

    haystack = " ".join(
        clean_value(row.get(field))
        for field in ("Address1", "Address2", "Address3", "Address4", "PlaceName")
    ).upper()
    for state in STATE_NAMES:
        if state.upper() in haystack:
            return state
    return ""


def infer_city(row: dict[str, str], state: str) -> str:
    place_name = parse_place_name(row.get("PlaceName"))
    if place_name and clean_key(place_name) != clean_key(state):
        return place_name

    for field in ("Address4", "Address3", "Address2", "Address1"):
        value = clean_value(row.get(field))
        if not value:
            continue
        candidate = re.sub(r"\b\d{6}\b", "", value).strip(" ,.-")
        if candidate and clean_key(candidate) != clean_key(state):
            return candidate

    return state or "Unknown"


def build_address_fields(row: dict[str, str]) -> dict[str, str] | None:
    address_parts = [tidy_address_part(row.get(f"Address{i}")) for i in range(1, 5)]
    address_parts = [part for part in address_parts if part]
    gst_no = clean_value(row.get("GSTNo"))
    has_address = bool(address_parts)
    has_gst = bool(gst_no)

    if not (has_address or has_gst):
        return None

    state = infer_state(row)
    city = infer_city(row, state)
    place_name = parse_place_name(row.get("PlaceName"))

    address_line1 = address_parts[0] if address_parts else place_name or clean_value(row.get("PartyName"))
    address_line2 = ", ".join(address_parts[1:]) if len(address_parts) > 1 else ""

    return {
        "address_title": clean_value(row.get("PartyName")),
        "address_type": "Billing",
        "address_line1": address_line1 or clean_value(row.get("PartyName")),
        "address_line2": address_line2,
        "city": city,
        "state": state,
        "country": "India",
        "pincode": extract_pincode(row),
        "is_primary_address": 1,
        "is_shipping_address": 1,
        "gstin": gst_no,
        "gst_category": "Registered Regular" if gst_no else "Unregistered",
    }


def build_details_text(row: dict[str, str]) -> str:
    lines = [f"Source Party Role: {clean_value(row.get('PartyRole'))}"]
    if clean_value(row.get("PartyAlias")):
        lines.append(f"Source Party Alias: {clean_value(row.get('PartyAlias'))}")
    if clean_value(row.get("AccCode")):
        lines.append(f"Source AccCode: {clean_value(row.get('AccCode'))}")
    if clean_value(row.get("GSTNo")):
        lines.append(f"Source GSTNo: {clean_value(row.get('GSTNo'))}")
    if clean_value(row.get("PANNo")):
        lines.append(f"Source PANNo: {clean_value(row.get('PANNo'))}")
    return "\n".join(lines)


def get_target_bucket(role: str) -> str:
    if role in CUSTOMER_ROLES:
        return "customer"
    if role in SUPPLIER_ROLES:
        return "supplier"
    if role in DRIVER_ROLES:
        return "driver"
    return "unmapped"


def get_unmapped_reason(role: str) -> tuple[str, str]:
    if role in {"Vehicle", "Tanker", "Vehicle Manufacturer"}:
        return ("Vehicle", "Source row does not include the required ERPNext Vehicle fields.")
    if role in {"Employ"}:
        return ("Employee", "This looks like employee data and should be reviewed before import.")
    if role in {"Vendor", "STORE", "Buyer", "BRANCH", "Company"}:
        return ("Manual Review", "Role mixes party and ledger-style records, so it was not auto-imported.")
    if role in ACCOUNT_LIKE_ROLES:
        return ("Account", "This role looks like a ledger or account head rather than a party master.")
    if role in {"Operator"}:
        return ("Manual Review", "This row is incomplete and needs manual cleanup.")
    return ("Manual Review", "No safe default ERPNext master mapping was inferred.")


def write_csv(path: Path, headers: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in headers})


def write_xlsx(path: Path, sheet_name: str, headers: list[str], rows: list[dict[str, object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for index, header in enumerate(headers, start=1):
        max_length = len(header)
        for row_index in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row=row_index, column=index).value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 60)

    workbook.save(path)


def write_dataset(base_path: Path, sheet_name: str, headers: list[str], rows: list[dict[str, object]]) -> None:
    write_csv(base_path.with_suffix(".csv"), headers, rows)
    write_xlsx(base_path.with_suffix(".xlsx"), sheet_name, headers, rows)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(SOURCE_FILE, data_only=True)
    worksheet = workbook.active
    headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
    source_rows = []
    for row_number in range(2, worksheet.max_row + 1):
        row = {header: worksheet.cell(row_number, headers.index(header) + 1).value for header in headers}
        row = {key: clean_value(value) for key, value in row.items()}
        row["_source_row"] = row_number
        source_rows.append(row)

    today_stamp = date.today().strftime("%Y%m%d")
    counters = {"customer": 0, "supplier": 0, "driver": 0}

    customer_rows: list[dict[str, object]] = []
    supplier_rows: list[dict[str, object]] = []
    driver_rows: list[dict[str, object]] = []
    address_rows: list[dict[str, object]] = []
    name_map_rows: list[dict[str, object]] = []
    unmapped_rows: list[dict[str, object]] = []

    for row in source_rows:
        role = clean_value(row.get("PartyRole"))
        bucket = get_target_bucket(role)
        source_party_name = clean_value(row.get("PartyName"))
        party_type = infer_party_type(row.get("PANNo"), source_party_name, role)

        if bucket == "customer":
            counters["customer"] += 1
            target_name = f"IMP-{today_stamp}-CUST-{counters['customer']:04d}"
            customer_rows.append(
                {
                    "name": target_name,
                    "customer_name": source_party_name,
                    "customer_type": party_type,
                    "customer_group": "All Customer Groups",
                    "territory": "All Territories",
                    "tax_id": clean_value(row.get("PANNo")),
                    "customer_details": build_details_text(row),
                }
            )
            name_map_rows.append(
                {
                    "source_row": row["_source_row"],
                    "party_role": role,
                    "source_party_name": source_party_name,
                    "source_party_alias": clean_value(row.get("PartyAlias")),
                    "source_acc_code": clean_value(row.get("AccCode")),
                    "target_doctype": "Customer",
                    "target_name": target_name,
                    "target_display_name": source_party_name,
                }
            )
            if address := build_address_fields(row):
                address_rows.append(
                    {
                        **address,
                        "links.link_doctype": "Customer",
                        "links.link_name": target_name,
                    }
                )
            continue

        if bucket == "supplier":
            counters["supplier"] += 1
            target_name = f"IMP-{today_stamp}-SUP-{counters['supplier']:04d}"
            supplier_rows.append(
                {
                    "name": target_name,
                    "supplier_name": source_party_name,
                    "supplier_type": party_type,
                    "supplier_group": "All Supplier Groups",
                    "country": "India",
                    "tax_id": clean_value(row.get("PANNo")),
                    "is_transporter": infer_is_transporter(role, source_party_name),
                    "supplier_details": build_details_text(row),
                }
            )
            name_map_rows.append(
                {
                    "source_row": row["_source_row"],
                    "party_role": role,
                    "source_party_name": source_party_name,
                    "source_party_alias": clean_value(row.get("PartyAlias")),
                    "source_acc_code": clean_value(row.get("AccCode")),
                    "target_doctype": "Supplier",
                    "target_name": target_name,
                    "target_display_name": source_party_name,
                }
            )
            if address := build_address_fields(row):
                address_rows.append(
                    {
                        **address,
                        "links.link_doctype": "Supplier",
                        "links.link_name": target_name,
                    }
                )
            continue

        if bucket == "driver":
            counters["driver"] += 1
            target_name = f"IMP-{today_stamp}-DRI-{counters['driver']:04d}"
            driver_rows.append(
                {
                    "name": target_name,
                    "full_name": source_party_name,
                    "status": "Active",
                }
            )
            name_map_rows.append(
                {
                    "source_row": row["_source_row"],
                    "party_role": role,
                    "source_party_name": source_party_name,
                    "source_party_alias": clean_value(row.get("PartyAlias")),
                    "source_acc_code": clean_value(row.get("AccCode")),
                    "target_doctype": "Driver",
                    "target_name": target_name,
                    "target_display_name": source_party_name,
                }
            )
            if address := build_address_fields(row):
                address_rows.append(
                    {
                        **address,
                        "links.link_doctype": "Driver",
                        "links.link_name": target_name,
                    }
                )
            continue

        suggested_doctype, reason = get_unmapped_reason(role)
        unmapped_rows.append(
            {
                "source_row": row["_source_row"],
                "suggested_doctype": suggested_doctype,
                "reason": reason,
                "PartyName": source_party_name,
                "PartyRole": role,
                "PartyAlias": clean_value(row.get("PartyAlias")),
                "Address1": clean_value(row.get("Address1")),
                "Address2": clean_value(row.get("Address2")),
                "Address3": clean_value(row.get("Address3")),
                "Address4": clean_value(row.get("Address4")),
                "PlaceName": clean_value(row.get("PlaceName")),
                "GSTNo": clean_value(row.get("GSTNo")),
                "PANNo": clean_value(row.get("PANNo")),
                "AccCode": clean_value(row.get("AccCode")),
            }
        )

    customer_headers = [
        "name",
        "customer_name",
        "customer_type",
        "customer_group",
        "territory",
        "tax_id",
        "customer_details",
    ]
    supplier_headers = [
        "name",
        "supplier_name",
        "supplier_type",
        "supplier_group",
        "country",
        "tax_id",
        "is_transporter",
        "supplier_details",
    ]
    driver_headers = ["name", "full_name", "status"]
    address_headers = [
        "address_title",
        "address_type",
        "address_line1",
        "address_line2",
        "city",
        "state",
        "country",
        "pincode",
        "is_primary_address",
        "is_shipping_address",
        "gstin",
        "gst_category",
        "links.link_doctype",
        "links.link_name",
    ]
    name_map_headers = [
        "source_row",
        "party_role",
        "source_party_name",
        "source_party_alias",
        "source_acc_code",
        "target_doctype",
        "target_name",
        "target_display_name",
    ]
    unmapped_headers = [
        "source_row",
        "suggested_doctype",
        "reason",
        "PartyName",
        "PartyRole",
        "PartyAlias",
        "Address1",
        "Address2",
        "Address3",
        "Address4",
        "PlaceName",
        "GSTNo",
        "PANNo",
        "AccCode",
    ]

    write_dataset(OUTPUT_DIR / "customer_import_v15", "Customers", customer_headers, customer_rows)
    write_dataset(OUTPUT_DIR / "supplier_import_v15", "Suppliers", supplier_headers, supplier_rows)
    write_dataset(OUTPUT_DIR / "driver_import_v15", "Drivers", driver_headers, driver_rows)
    write_dataset(OUTPUT_DIR / "address_import_v15", "Addresses", address_headers, address_rows)
    write_dataset(OUTPUT_DIR / "import_name_map", "Name Map", name_map_headers, name_map_rows)
    write_dataset(OUTPUT_DIR / "unmapped_review_v15", "Unmapped", unmapped_headers, unmapped_rows)

    summary_lines = [
        "# ERPNext v15 Party Import Outputs",
        "",
        f"Source file: `{SOURCE_FILE}`",
        f"Generated on: `{date.today().isoformat()}`",
        "",
        "Import order:",
        "1. `customer_import_v15.xlsx` or `.csv`",
        "2. `supplier_import_v15.xlsx` or `.csv`",
        "3. `driver_import_v15.xlsx` or `.csv`",
        "4. `address_import_v15.xlsx` or `.csv`",
        "",
        "Counts:",
        f"- Customers: {len(customer_rows)}",
        f"- Suppliers: {len(supplier_rows)}",
        f"- Drivers: {len(driver_rows)}",
        f"- Addresses: {len(address_rows)}",
        f"- Unmapped review rows: {len(unmapped_rows)}",
        "",
        "Mapping assumptions:",
        "- `CLIENT` rows were mapped to `Customer`.",
        "- `Supplier`, `Tyre Supplier`, `PUMP`, `Tax Vendor`, `OWNER`, `Broker`, `Agent`, `Financier`, and `Bank` rows were mapped to `Supplier`.",
        "- `Driver Ledger` and `Debit Driver` rows were mapped to `Driver`.",
        "- Explicit `name` values were assigned so the address file can link to the imported masters reliably.",
        "- Address rows were generated only when the source row had at least one address line or a GST number.",
        "- Remaining roles were kept in `unmapped_review_v15` for manual review because they look like ledger heads, vehicles, or mixed-purpose records.",
    ]
    (OUTPUT_DIR / "README.md").write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    print(f"Generated files in {OUTPUT_DIR}")
    print(
        "Summary:",
        {
            "customers": len(customer_rows),
            "suppliers": len(supplier_rows),
            "drivers": len(driver_rows),
            "addresses": len(address_rows),
            "unmapped": len(unmapped_rows),
        },
    )


if __name__ == "__main__":
    main()
